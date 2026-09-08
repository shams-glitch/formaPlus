# -*- coding: utf-8 -*-
"""Génère la présentation jury GCentreFormation — FormaPlus (10 slides)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import nsmap
from pptx.oxml import parse_xml
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
CAPS = DOCS / "presentation_captures"
OUT_PPTX = DOCS / "GCentreFormation_Presentation_Septembre_2026.pptx"

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
NAVY2 = RGBColor(0x12, 0x33, 0x58)
TEAL = RGBColor(0x1A, 0xA3, 0x96)
TEAL_DK = RGBColor(0x0E, 0x7A, 0x72)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFFWHITE = RGBColor(0xF4, 0xF7, 0xFA)
SLATE = RGBColor(0x4A, 0x5B, 0x6B)
INK = RGBColor(0x1C, 0x2B, 0x3A)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD5, 0xDD, 0xE5)
RED = RGBColor(0xC0, 0x39, 0x2B)

W = Inches(13.333)
H = Inches(7.5)


def set_run(run, text, size=18, bold=False, color=INK, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_text_box(slide, l, t, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    set_run(run, text, size, bold, color)
    return box


def box_lines(slide, l, t, w, h, lines, size=16, color=INK, bold_first=False, spacing=8):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        run = p.add_run()
        set_run(run, line, size, bold=(bold_first and i == 0), color=color)
    return box


def rect(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
    return sh


def round_rect(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
    return sh


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def footer(slide, num):
    add_text_box(slide, Inches(0.5), Inches(7.18), Inches(10), Inches(0.28),
                 "GCentreFormation  ·  FormaPlus  ·  Validation septembre 2026",
                 11, False, SLATE)
    add_text_box(slide, Inches(11.6), Inches(7.18), Inches(1.2), Inches(0.28),
                 str(num), 11, False, SLATE, PP_ALIGN.RIGHT)


def title_bar(slide, title, subtitle=None):
    rect(slide, 0, 0, W, Inches(0.12), TEAL)
    rect(slide, 0, Inches(0.12), W, Inches(1.05), NAVY)
    add_text_box(slide, Inches(0.5), Inches(0.22), Inches(12), Inches(0.5),
                 title, 26, True, WHITE)
    if subtitle:
        add_text_box(slide, Inches(0.5), Inches(0.70), Inches(12), Inches(0.35),
                     subtitle, 13, False, RGBColor(0xB8, 0xC9, 0xD6))


def add_picture_fit(slide, path, l, t, w, h):
    pic = slide.shapes.add_picture(str(path), l, t, w, h)
    return pic


def slide1(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, NAVY)
    rect(s, 0, 0, Inches(0.18), H, TEAL)
    add_text_box(s, Inches(0.8), Inches(1.5), Inches(11), Inches(0.4),
                 "VALIDATION  ·  SEPTEMBRE 2026", 14, True, TEAL)
    add_text_box(s, Inches(0.8), Inches(2.0), Inches(12), Inches(0.85),
                 "GCentreFormation", 44, True, WHITE)
    add_text_box(s, Inches(0.8), Inches(2.85), Inches(12), Inches(0.55),
                 "FormaPlus", 32, True, TEAL)
    add_text_box(s, Inches(0.8), Inches(3.6), Inches(11.5), Inches(0.5),
                 "Application desktop de gestion d’un centre de formation",
                 20, False, RGBColor(0xD5, 0xDE, 0xE6))
    rect(s, Inches(0.8), Inches(4.35), Inches(2.2), Inches(0.06), TEAL)
    add_text_box(s, Inches(0.8), Inches(4.7), Inches(11), Inches(0.4),
                 "Shamsedine Tayeb", 18, True, WHITE)
    add_text_box(s, Inches(0.8), Inches(5.15), Inches(11), Inches(0.4),
                 "Qt / C++  ·  Oracle XE  ·  Modules Formateurs et Cours",
                 15, False, RGBColor(0xA8, 0xBB, 0xC8))
    notes(s, (
        "Bonjour. Je m’appelle Shamsedine Tayeb. Je vous présente GCentreFormation, "
        "aussi appelé FormaPlus. C’est une application desktop écrite en C++, avec Qt 6.8.3 et qmake. "
        "Elle sert à gérer un centre de formation. J’ai choisi deux modules parmi ceux du sujet : "
        "les formateurs et les cours. La base de données est Oracle XE, reliée par QODBC. "
        "Dans les minutes qui viennent, je vais vous montrer le périmètre, l’architecture réelle, "
        "la base, le CRUD, puis les métiers qui sont vraiment dans le code. Ensuite je ferai la démo."
    ))


def slide2(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, OFFWHITE)
    title_bar(s, "Contexte et objectif", "Pourquoi cette application")
    items = [
        ("Centre de formation", "Besoin de suivre les formateurs et le catalogue de cours au même endroit."),
        ("Deux modules liés", "Un cours a un formateur. Un formateur peut n’avoir aucun cours, ou plusieurs."),
        ("Objectif", "Une application desktop claire pour le CRUD, le suivi et les contrôles."),
        ("Périmètre officiel", "Sujet de repassage septembre 2026 — uniquement les exigences de l’énoncé."),
    ]
    for i, (t, d) in enumerate(items):
        x = Inches(0.5 + (i % 2) * 6.4)
        y = Inches(1.55 + (i // 2) * 2.5)
        round_rect(s, x, y, Inches(6.1), Inches(2.2), WHITE, LINE)
        rect(s, x, y, Inches(0.12), Inches(2.2), TEAL)
        add_text_box(s, x + Inches(0.4), y + Inches(0.3), Inches(5.5), Inches(0.45), t, 20, True, NAVY)
        add_text_box(s, x + Inches(0.4), y + Inches(0.9), Inches(5.5), Inches(1.0), d, 16, False, SLATE)
    footer(s, 2)
    notes(s, (
        "Le besoin est concret. Un centre de formation doit suivre ses formateurs et son catalogue "
        "de cours au même endroit. Ces deux modules sont liés : un cours a toujours un formateur, "
        "un formateur peut n’avoir aucun cours ou plusieurs. L’objectif, c’est une application desktop "
        "claire, pas un site web. On peut ajouter, modifier, supprimer, contrôler les données, "
        "et voir les indicateurs. Je me suis limité au sujet officiel de septembre 2026. "
        "Je n’ai pas ajouté de fonctionnalités hors énoncé."
    ))


def slide3(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, OFFWHITE)
    title_bar(s, "Modules choisis", "Formateurs  +  Cours")
    # left
    round_rect(s, Inches(0.45), Inches(1.5), Inches(6.1), Inches(5.4), WHITE, LINE)
    rect(s, Inches(0.45), Inches(1.5), Inches(6.1), Inches(0.7), NAVY)
    add_text_box(s, Inches(0.7), Inches(1.62), Inches(5.6), Inches(0.5), "FORMATEURS", 22, True, WHITE)
    for i, line in enumerate([
        "CRUD complet",
        "Recherche / tri  ≥ 3 critères",
        "Statistiques + graphiques live",
        "PDF fiche et liste",
        "Charge pédagogique",
        "Conflits de planning",
    ]):
        y = Inches(2.45 + i * 0.65)
        round_rect(s, Inches(0.75), y, Inches(5.5), Inches(0.52), RGBColor(0xE8, 0xF6, 0xF4), None)
        add_text_box(s, Inches(0.95), y + Inches(0.08), Inches(5.1), Inches(0.38), "▸  " + line, 16, False, NAVY)
    # right
    round_rect(s, Inches(6.8), Inches(1.5), Inches(6.1), Inches(5.4), WHITE, LINE)
    rect(s, Inches(6.8), Inches(1.5), Inches(6.1), Inches(0.7), TEAL_DK)
    add_text_box(s, Inches(7.05), Inches(1.62), Inches(5.6), Inches(0.5), "COURS", 22, True, WHITE)
    for i, line in enumerate([
        "CRUD complet",
        "Recherche / tri  ≥ 3 critères",
        "Statistiques + graphiques live",
        "PDF fiche et catalogue",
        "Affectation formateur",
        "Alertes métier",
    ]):
        y = Inches(2.45 + i * 0.65)
        round_rect(s, Inches(7.1), y, Inches(5.5), Inches(0.52), RGBColor(0xEA, 0xF2, 0xF8), None)
        add_text_box(s, Inches(7.3), y + Inches(0.08), Inches(5.1), Inches(0.38), "▸  " + line, 16, False, NAVY)
    footer(s, 3)
    notes(s, (
        "Le sujet demande deux modules. J’ai choisi Formateurs et Cours, parce qu’ils sont liés. "
        "Pour chaque module, le CRUD est implémenté. Il y a aussi la recherche et le tri avec au moins "
        "trois critères, les statistiques avec graphiques, et un document PDF personnalisé. "
        "Ensuite deux métiers utiles. Côté formateurs : la charge pédagogique et les conflits de planning. "
        "Côté cours : l’affectation d’un formateur et les alertes. Tout ça est dans l’application, "
        "pas seulement dans le rapport. Je pourrai le montrer à l’écran."
    ))


def slide4(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, OFFWHITE)
    title_bar(s, "Architecture", "Séparation GUI / accès données — pas de SQL derrière un bouton")
    layers = [
        ("UTILISATEUR", NAVY),
        ("QT GUI   ·   Sidebar + StackedWidget + Tabs", NAVY2),
        ("Formateur   /   Cours    (classes C++ pures)", TEAL_DK),
        ("Connection::instance()    Singleton", TEAL),
        ("QODBC", NAVY2),
        ("ORACLE XE    DSN Source_Projet2A", NAVY),
    ]
    y = Inches(1.45)
    for label, color in layers:
        round_rect(s, Inches(2.4), y, Inches(8.5), Inches(0.62), color)
        add_text_box(s, Inches(2.4), y + Inches(0.12), Inches(8.5), Inches(0.42),
                     label, 16, True, WHITE, PP_ALIGN.CENTER)
        y += Inches(0.72)
        if label != layers[-1][0]:
            add_text_box(s, Inches(6.3), y - Inches(0.16), Inches(0.8), Inches(0.2),
                         "↓", 14, True, TEAL, PP_ALIGN.CENTER)
    pts = [
        "GUI séparée des requêtes",
        "SQL dans Formateur et Cours",
        "Une seule connexion Oracle",
        "qmake  ·  Qt 6.8.3  ·  C++17",
    ]
    for i, p in enumerate(pts):
        add_text_box(s, Inches(0.4), Inches(1.55 + i * 1.2), Inches(1.9), Inches(1.0),
                     p, 13, False, SLATE, PP_ALIGN.CENTER)
    footer(s, 4)
    notes(s, (
        "Voici l’architecture réelle. L’utilisateur passe par l’interface Qt : une barre latérale, "
        "un Stacked Widget et des onglets. Les boutons n’écrivent pas de SQL. Ils appellent les classes "
        "Formateur et Cours. C’est là que se trouvent toutes les requêtes. La connexion passe par un "
        "Singleton, Connection::instance. Ensuite QODBC, puis Oracle XE, avec le DSN Source_Projet2A. "
        "Le sujet impose une seule connexion. Le projet est compilé avec qmake, Qt 6.8.3 et MinGW. "
        "Je ne présente pas ça comme un MVC : ce n’est pas le patron que j’ai nommé dans le code."
    ))


def slide5(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, OFFWHITE)
    title_bar(s, "Base de données", "MCD → MLD → MPD → Oracle")
    steps = ["MCD", "MLD", "MPD", "Oracle"]
    for i, name in enumerate(steps):
        x = Inches(0.5 + i * 3.2)
        round_rect(s, x, Inches(1.5), Inches(2.7), Inches(0.7), NAVY if i < 3 else TEAL)
        add_text_box(s, x, Inches(1.6), Inches(2.7), Inches(0.5), name, 20, True, WHITE, PP_ALIGN.CENTER)
        if i < 3:
            add_text_box(s, x + Inches(2.55), Inches(1.58), Inches(0.7), Inches(0.5),
                         "→", 20, True, TEAL, PP_ALIGN.CENTER)
    # relationship
    round_rect(s, Inches(0.5), Inches(2.5), Inches(6.3), Inches(4.3), WHITE, LINE)
    add_text_box(s, Inches(0.75), Inches(2.65), Inches(5.8), Inches(0.4),
                 "Association  enseigne", 16, True, NAVY)
    round_rect(s, Inches(1.0), Inches(3.3), Inches(2.4), Inches(2.6), NAVY)
    add_text_box(s, Inches(1.0), Inches(3.7), Inches(2.4), Inches(0.5), "FORMATEUR", 16, True, WHITE, PP_ALIGN.CENTER)
    add_text_box(s, Inches(1.0), Inches(4.3), Inches(2.4), Inches(0.8), "PK  id_formateur\n(0,N)", 14, False, RGBColor(0xC8, 0xDC, 0xE8), PP_ALIGN.CENTER)
    add_text_box(s, Inches(3.45), Inches(4.2), Inches(1.3), Inches(0.5), "enseigne", 12, True, TEAL, PP_ALIGN.CENTER)
    round_rect(s, Inches(4.7), Inches(3.3), Inches(1.8), Inches(2.6), TEAL_DK)
    add_text_box(s, Inches(4.7), Inches(3.7), Inches(1.8), Inches(0.5), "COURS", 16, True, WHITE, PP_ALIGN.CENTER)
    add_text_box(s, Inches(4.7), Inches(4.3), Inches(1.8), Inches(0.8), "PK  id_cours\n(1,1)", 14, False, WHITE, PP_ALIGN.CENTER)

    cons = [
        ("PK", "ID_FORMATEUR, ID_COURS"),
        ("FK", "COURS.ID_FORMATEUR → FORMATEUR"),
        ("UNIQUE", "CIN, e-mail"),
        ("CHECK", "CIN, statuts, dates, prix"),
        ("Séquences / triggers", "SEQ_*  +  TRG_*_BI"),
        ("Intégrité", "ORA-02292 si formateur utilisé"),
    ]
    for i, (k, v) in enumerate(cons):
        y = Inches(2.5 + i * 0.7)
        round_rect(s, Inches(7.05), y, Inches(5.8), Inches(0.62), WHITE, LINE)
        add_text_box(s, Inches(7.2), y + Inches(0.12), Inches(2.0), Inches(0.4), k, 14, True, TEAL)
        add_text_box(s, Inches(9.2), y + Inches(0.12), Inches(3.5), Inches(0.4), v, 13, False, INK)
    footer(s, 5)
    notes(s, (
        "La conception suit MCD, puis MLD, puis MPD Oracle. L’association s’appelle enseigne. "
        "Un formateur enseigne zéro à N cours. Un cours a exactement un formateur. "
        "Au niveau MCD, on ne met pas encore la clé étrangère. Elle apparaît au MLD, portée par COURS. "
        "Dans Oracle, on a les clés primaires, la clé étrangère, UNIQUE sur le CIN et l’e-mail, "
        "des CHECK, plus des séquences et des triggers pour générer les ID. "
        "Si on essaie de supprimer un formateur qui a encore des cours, Oracle refuse avec ORA-02292. "
        "Il n’y a pas de ON DELETE CASCADE."
    ))


def slide6(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, OFFWHITE)
    title_bar(s, "CRUD + validation", "Captures réelles de l’application")
    cap_f = CAPS / "02_formateurs_crud.png"
    cap_c = CAPS / "04_cours_crud.png"
    add_picture_fit(s, cap_f, Inches(0.35), Inches(1.4), Inches(6.3), Inches(3.55))
    add_picture_fit(s, cap_c, Inches(6.75), Inches(1.4), Inches(6.3), Inches(3.55))
    add_text_box(s, Inches(0.35), Inches(4.95), Inches(6.3), Inches(0.3),
                 "Formateurs  ·  CREATE READ UPDATE DELETE", 12, True, NAVY, PP_ALIGN.CENTER)
    add_text_box(s, Inches(6.75), Inches(4.95), Inches(6.3), Inches(0.3),
                 "Cours  ·  CREATE READ UPDATE DELETE", 12, True, NAVY, PP_ALIGN.CENTER)
    chips = [
        "valider() C++",
        "CHECK / UNIQUE / FK Oracle",
        "Messages d’erreur métier",
        "Intégrité référentielle",
    ]
    for i, c in enumerate(chips):
        x = Inches(0.4 + i * 3.2)
        round_rect(s, x, Inches(5.4), Inches(3.05), Inches(0.7), NAVY)
        add_text_box(s, x, Inches(5.52), Inches(3.05), Inches(0.45), c, 13, True, WHITE, PP_ALIGN.CENTER)
    footer(s, 6)
    notes(s, (
        "Voici les écrans réels de l’application, pas des maquettes. "
        "À gauche, les formateurs. À droite, les cours. Dans les deux cas : ajouter, afficher, "
        "modifier, supprimer. Avant l’INSERT ou l’UPDATE, la méthode valider contrôle la saisie. "
        "Par exemple, le CIN doit avoir exactement 8 chiffres, l’e-mail doit être valide. "
        "Oracle contrôle aussi, avec CHECK, UNIQUE et la clé étrangère. "
        "Si un formateur a encore des cours, la suppression est refusée. "
        "L’utilisateur voit un message métier, pas seulement un code Oracle."
    ))


def slide7(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, OFFWHITE)
    title_bar(s, "Métiers implémentés", "Fonctionnalités avancées réellement codées")
    common = [
        ("Recherche ≥ 3", "Nom, spécialité, statut\nTitre, niveau, statut"),
        ("Tri", "Liste blanche de colonnes\nASC / DESC"),
        ("Stats live", "chargerTout()\naprès chaque CRUD"),
        ("PDF", "QPdfWriter personnalisé\nfiche / liste / catalogue"),
    ]
    for i, (t, d) in enumerate(common):
        x = Inches(0.4 + i * 3.2)
        round_rect(s, x, Inches(1.45), Inches(3.05), Inches(2.15), WHITE, LINE)
        rect(s, x, Inches(1.45), Inches(3.05), Inches(0.12), TEAL)
        add_text_box(s, x + Inches(0.15), Inches(1.7), Inches(2.75), Inches(0.4), t, 16, True, NAVY)
        add_text_box(s, x + Inches(0.15), Inches(2.2), Inches(2.75), Inches(1.15), d, 13, False, SLATE)
    extras = [
        ("Formateurs", TEAL, ["Charge pédagogique  —  cours + heures", "Conflits  —  mêmes dates, même formateur"]),
        ("Cours", NAVY, ["Affectation  —  UPDATE ID_FORMATEUR", "Alertes  —  7 jours, inactif, chevauchement"]),
    ]
    for i, (t, col, lines) in enumerate(extras):
        x = Inches(0.4 + i * 6.45)
        round_rect(s, x, Inches(3.85), Inches(6.25), Inches(2.55), WHITE, LINE)
        rect(s, x, Inches(3.85), Inches(6.25), Inches(0.55), col)
        add_text_box(s, x + Inches(0.25), Inches(3.93), Inches(5.8), Inches(0.4), t, 16, True, WHITE)
        for j, line in enumerate(lines):
            add_text_box(s, x + Inches(0.3), Inches(4.6 + j * 0.7), Inches(5.7), Inches(0.55),
                         "▸  " + line, 15, False, INK)
    footer(s, 7)
    notes(s, (
        "Les métiers communs aux deux modules : recherche multicritères, au moins trois critères, "
        "tri, statistiques, et graphiques mis à jour après chaque CRUD grâce à chargerTout. "
        "Le document est uniquement en PDF, généré avec QPdfWriter. Ce n’est pas une capture d’écran : "
        "il y a un en-tête FormaPlus, des tableaux et un pied de page. "
        "En plus, pour les formateurs, je calcule la charge pédagogique et je détecte les conflits "
        "quand le même formateur a des cours qui se chevauchent. Pour les cours, j’affecte un formateur "
        "et j’affiche des alertes : début dans sept jours, formateur inactif, ou chevauchement."
    ))


def slide8(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, OFFWHITE)
    title_bar(s, "Qualité et accès aux données", "Pratiques réellement présentes dans le code")
    cards = [
        ("1  Singleton", "Connection::instance()\nconstructeur privé\ncopie = delete"),
        ("2  Requêtes préparées", "prepare()\nbindValue()\nexec()"),
        ("3  Double validation", "valider() en C++\nCHECK / UNIQUE / FK"),
        ("4  SQL hors GUI", "pas de SQL dans les slots\nclasses Formateur / Cours"),
    ]
    for i, (t, d) in enumerate(cards):
        x = Inches(0.4 + i * 3.2)
        round_rect(s, x, Inches(1.5), Inches(3.05), Inches(3.5), WHITE, LINE)
        rect(s, x, Inches(1.5), Inches(3.05), Inches(0.7), NAVY if i % 2 == 0 else TEAL_DK)
        add_text_box(s, x + Inches(0.15), Inches(1.62), Inches(2.75), Inches(0.5), t, 16, True, WHITE)
        add_text_box(s, x + Inches(0.2), Inches(2.45), Inches(2.65), Inches(2.2), d, 16, False, INK)
    round_rect(s, Inches(0.4), Inches(5.25), Inches(12.5), Inches(1.15), NAVY)
    add_text_box(s, Inches(0.65), Inches(5.5), Inches(12.1), Inches(0.7),
                 "Les requêtes préparées réduisent les risques d’injection SQL.",
                 20, True, WHITE, PP_ALIGN.CENTER)
    footer(s, 8)
    notes(s, (
        "Quatre points présents dans le code. Premier : le Singleton. Connection::instance crée "
        "une seule connexion. Le constructeur est privé, la copie est interdite. "
        "Deuxième : prepare, bindValue, exec. On n’assemble pas les valeurs dans la chaîne SQL. "
        "Les requêtes préparées réduisent les risques d’injection SQL. Pour le tri, seules des colonnes "
        "autorisées sont acceptées. Troisième : validation des deux côtés, C++ et Oracle. "
        "Quatrième : aucune requête derrière un bouton. Les slots appellent seulement le modèle. "
        "Je ne revendique pas d’autre mécanisme de sécurité que ça."
    ))


def slide9(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, OFFWHITE)
    title_bar(s, "Démonstration de l’application", "Ordre de la démo — captures réelles")
    add_picture_fit(s, CAPS / "01_accueil.png", Inches(4.7), Inches(1.4), Inches(8.2), Inches(5.35))
    steps = [
        "1  Accueil  KPI + graphiques",
        "2  Formateurs CRUD",
        "3  Recherche / tri",
        "4  Statistiques live",
        "5  PDF fiche / liste",
        "6  Charge / conflits",
        "7  Cours CRUD",
        "8  Recherche / stats",
        "9  Affectation / alertes",
    ]
    for i, line in enumerate(steps):
        y = Inches(1.4 + i * 0.58)
        fill = TEAL if i == 0 else WHITE
        col = WHITE if i == 0 else INK
        round_rect(s, Inches(0.35), y, Inches(4.15), Inches(0.52), fill, None if i == 0 else LINE)
        add_text_box(s, Inches(0.5), y + Inches(0.06), Inches(3.9), Inches(0.4), line, 13, True if i == 0 else False, col)
    footer(s, 9)
    notes(s, (
        "Je passe à la démonstration, dans cet ordre. D’abord l’accueil : les indicateurs et les graphiques. "
        "Ensuite les formateurs : CRUD, puis recherche et tri, puis les statistiques qui se mettent à jour, "
        "puis le PDF, puis la charge et les conflits. Ensuite les cours : CRUD, recherche et statistiques, "
        "affectation et alertes. Si le jury le demande, je tente de supprimer un formateur qui a encore "
        "des cours : Oracle refuse. Je ne supprime pas de vraies données importantes juste pour montrer. "
        "L’image à droite est une capture réelle de l’accueil."
    ))


def slide10(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, NAVY)
    rect(s, 0, 0, Inches(0.18), H, TEAL)
    add_text_box(s, Inches(0.7), Inches(0.55), Inches(12), Inches(0.5),
                 "Conclusion", 28, True, WHITE)
    points = [
        "Deux modules implémentés  —  Formateurs et Cours",
        "CRUD opérationnel, avec validation C++ et Oracle",
        "Métiers : recherche, tri, stats live, PDF, charge, conflits, affectation, alertes",
        "Oracle intégré  —  PK, FK, CHECK, séquences, triggers",
        "Application desktop Qt / C++  —  une connexion Singleton, requêtes préparées",
        "Projet prêt pour la validation de septembre 2026",
    ]
    for i, p in enumerate(points):
        add_text_box(s, Inches(0.85), Inches(1.25 + i * 0.55), Inches(11.8), Inches(0.5),
                     "▸   " + p, 16, False, RGBColor(0xD5, 0xDE, 0xE6))
    rect(s, Inches(0.85), Inches(4.7), Inches(2.0), Inches(0.06), TEAL)
    add_text_box(s, Inches(0.85), Inches(5.05), Inches(11), Inches(0.5),
                 "Merci pour votre attention.", 26, True, WHITE)
    add_text_box(s, Inches(0.85), Inches(5.65), Inches(11), Inches(0.45),
                 "Questions ?", 22, True, TEAL)
    add_text_box(s, Inches(0.85), Inches(6.5), Inches(11), Inches(0.35),
                 "Shamsedine Tayeb  ·  GCentreFormation / FormaPlus", 13, False, RGBColor(0x9A, 0xB0, 0xBE))
    notes(s, (
        "Pour conclure, le projet implémente deux modules, Formateurs et Cours. "
        "Le CRUD fonctionne, avec une validation en C++ et dans Oracle. "
        "Les métiers demandés sont là : recherche, tri, statistiques live, PDF, plus la charge, "
        "les conflits, l’affectation et les alertes. Oracle est intégré avec PK, FK, CHECK, "
        "séquences et triggers. L’application est desktop Qt et C++, avec un Singleton et des "
        "requêtes préparées. Le projet est prêt pour la validation. Merci pour votre attention. "
        "Je suis disponible pour vos questions.\n\n"
        "Questions techniques possibles\n"
        "• Pourquoi Singleton ? Le sujet impose une seule connexion. Connection::instance crée une unique instance QODBC.\n"
        "• Pourquoi Oracle ? Imposé par l’énoncé. Ici Oracle XE via QODBC et le DSN Source_Projet2A.\n"
        "• Pourquoi Qt ? Le sujet demande une application desktop avec le framework Qt.\n"
        "• Où est le SQL ? Dans Formateur et Cours uniquement. Pas dans les slots des boutons.\n"
        "• Pourquoi prepare/bindValue ? La valeur n’est pas collée dans le SQL. Ça réduit l’injection SQL.\n"
        "• Comment fonctionne le CRUD ? ajouter, afficher, modifier, supprimer dans les classes métier, puis chargerTout.\n"
        "• Recherche multicritères ? Au moins trois critères, avec bindValue. Formateurs : nom, spécialité, statut. Cours : titre, niveau, statut.\n"
        "• Statistiques ? GROUP BY, puis Qt Charts. Mis à jour après chaque CRUD via chargerTout.\n"
        "• PDF ? QPdfWriter, document personnalisé, un seul format, pas une capture d’écran.\n"
        "• Relation Formateur/Cours ? enseigne : FORMATEUR (0,N) — COURS (1,1).\n"
        "• Contraintes FK ? Pas de CASCADE. Supprimer un formateur utilisé → ORA-02292.\n"
        "• Pourquoi Formateurs + Cours ? Deux modules liés, ça justifie la clé étrangère et les métiers.\n"
        "Arduino / marketing / teamwork : N/A — NOT REQUIRED BY SEPTEMBER 2026 SUBJECT."
    ))


def main():
    missing = [p for p in [
        CAPS / "01_accueil.png",
        CAPS / "02_formateurs_crud.png",
        CAPS / "04_cours_crud.png",
    ] if not p.exists()]
    if missing:
        raise SystemExit("Captures manquantes : " + ", ".join(str(p) for p in missing))

    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    slide1(prs)
    slide2(prs)
    slide3(prs)
    slide4(prs)
    slide5(prs)
    slide6(prs)
    slide7(prs)
    slide8(prs)
    slide9(prs)
    slide10(prs)
    prs.save(OUT_PPTX)
    print("WROTE", OUT_PPTX)


if __name__ == "__main__":
    main()
