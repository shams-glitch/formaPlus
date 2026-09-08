"""Generate Cahier de spécification Word document for GCentreFormation."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

OUT = Path(__file__).resolve().parent.parent / "docs" / "Cahier_Specification_GCentreFormation.docx"


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    p.paragraph_format.line_spacing = 1.5
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def main():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Cahier de spécification\n")
    r.bold = True
    r.font.size = Pt(18)
    r2 = title.add_run("GCentreFormation — FormaPlus\n")
    r2.font.size = Pt(16)
    title.add_run("Projet C++ 2A — Repassage Septembre 2026\n").font.size = Pt(14)
    title.add_run("Modules : Gestion des Formateurs & Gestion des Cours\n").font.size = Pt(12)
    title.add_run("Équipe FormaPlus\n").font.size = Pt(12)

    doc.add_page_break()
    add_heading(doc, "Table des matières", 1)
    add_para(doc, "(À actualiser dans Word : Références → Table des matières → Automatique)")

    add_heading(doc, "1. Contexte général", 1)
    add_heading(doc, "1.1 Étude de l'existant", 2)
    add_para(doc,
             "FormaPlus est un centre de formation professionnelle qui gère formateurs, "
             "catalogue de cours et planification pédagogique. Aujourd'hui, une partie "
             "des données peut être dispersée (fichiers, tableurs). Le client cible est "
             "le responsable pédagogique du centre, qui doit suivre les formateurs, "
             "les cours planifiés et détecter les conflits de planning.")

    add_heading(doc, "1.2 Problématique", 2)
    add_para(doc,
             "Comment centraliser la gestion des formateurs et des cours dans une "
             "application desktop fiable, avec persistance Oracle, contrôles d'intégrité "
             "et outils d'aide à la décision (statistiques, alertes, exports PDF) ?")

    add_heading(doc, "2. Solution proposée", 1)
    add_heading(doc, "2.1 Exigences fonctionnelles", 2)

    table = doc.add_table(rows=1, cols=5)
    hdr = table.rows[0].cells
    headers = ["Module", "Entité", "Attributs", "Utilisateur", "Fonctionnalités"]
    for i, h in enumerate(headers):
        hdr[i].text = h

    rows = [
        ("Gestion des Formateurs", "Formateur",
         "ID, CIN, nom, prénom, e-mail, téléphone, spécialité, date embauche, statut",
         "Responsable pédagogique",
         "CRUD ; recherche/tri (nom, spécialité, statut) ; stats graphiques ; PDF ; charge pédagogique ; conflits planning"),
        ("Gestion des Cours", "Cours",
         "ID, titre, niveau, durée, prix, dates, statut, ID formateur",
         "Responsable pédagogique",
         "CRUD ; recherche/tri (titre, niveau, statut, formateur) ; stats graphiques ; PDF catalogue ; affectation formateur ; alertes"),
    ]
    for mod, ent, attr, user, func in rows:
        row = table.add_row().cells
        row[0].text = mod
        row[1].text = ent
        row[2].text = attr
        row[3].text = user
        row[4].text = func

    add_para(doc, "Figure 1 — Tableau des exigences fonctionnelles par module", bold=True)

    add_heading(doc, "2.2 Exigences non fonctionnelles", 2)
    for item in [
        "SGBD Oracle obligatoire, une seule connexion (patron Singleton QODBC).",
        "Requêtes préparées dans les classes C++ Formateur et Cours (architecture Modèle-Vue).",
        "Contrôles de saisie C++ et contraintes CHECK/UNIQUE/FK en base.",
        "Interface Qt : un seul fichier .ui, Stacked Widget + Tab Widget, peu de boîtes de dialogue.",
        "Export document : PDF personnalisé uniquement (QPdfWriter, en-tête FormaPlus).",
        "Versionnement Git/GitHub avec commits réguliers.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "2.3 ODD touchés", 2)
    doc.add_paragraph("ODD 4 — Éducation de qualité : digitalisation de la gestion pédagogique.", style="List Bullet")
    doc.add_paragraph("ODD 8 — Travail décent : suivi de la charge des formateurs.", style="List Bullet")

    add_heading(doc, "3. Conception des données", 1)
    add_para(doc,
             "MCD : association enseigne entre FORMATEUR (0,N) et COURS (1,1). "
             "MLD : tables FORMATEUR et COURS avec clé étrangère ID_FORMATEUR. "
             "MPD Oracle : script sql/mpd_oracle.sql (séquences, triggers, CHECK ; "
             "FK sans ON DELETE = comportement restrict Oracle).")
    add_para(doc, "Figure 2 — Schéma MCD (voir docs/conception.md)", bold=True)

    add_heading(doc, "4. Conception graphique", 1)
    add_heading(doc, "4.1 Logo", 2)
    add_para(doc, "Logo textuel FormaPlus — barre latérale bleu nuit (#0F2744), accent vert (#1ABC9C).")

    add_heading(doc, "4.2 Charte graphique", 2)
    add_para(doc, "Police : Segoe UI / Arial. Couleurs : bleu nuit, vert accent, fond gris clair. "
             "Navigation : sidebar + pages empilées (Accueil, Formateurs, Cours).")

    add_heading(doc, "4.3 Interfaces GUI", 2)
    add_para(doc,
             "Captures d'écran à insérer avant la soutenance : tableau de bord KPI, "
             "onglets CRUD, recherche/tri multicritères, graphiques Qt Charts, "
             "onglets charge/conflits et alertes.")
    add_para(doc, "Figure 3 — Capture : tableau de bord Accueil (à insérer)", bold=True)
    add_para(doc, "Figure 4 — Capture : module Formateurs CRUD (à insérer)", bold=True)
    add_para(doc, "Figure 5 — Capture : module Cours — alertes (à insérer)", bold=True)

    add_heading(doc, "5. Architecture technique", 1)
    add_para(doc,
             "Couche connexion : Connection (Singleton). "
             "Couche modèle : Formateur, Cours. "
             "Couche vue : GCentreFormation (gcentreformation.ui). "
             "Aucune requête SQL dans les slots des boutons.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Generated: {OUT}")


if __name__ == "__main__":
    main()
